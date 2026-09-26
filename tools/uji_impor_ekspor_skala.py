"""
Debugging dengan cara berbeda: impor berantakan, ekspor, dan skala besar.

Tiga sudut yang belum pernah diserang sebelumnya.

Pertama, impor berkas. Berkas CSV yang datang dari aplikasi lain jarang
rapi: pemisahnya bisa titik koma atau koma, ada tanda kutip di dalam nama,
angka memakai titik sebagai pemisah ribuan, tanggal memakai format yang
berbeda beda, dan kadang ada baris kosong atau kolom yang bergeser. Impor
yang hanya benar untuk berkas yang rapi akan tampak berhasil tetapi
memasukkan angka yang salah, dan itu lebih berbahaya daripada gagal
dengan jelas.

Kedua, ekspor. Berkas yang dihasilkan diperiksa isinya, bukan hanya
apakah berkasnya terbentuk. Ekspor yang menghasilkan berkas kosong atau
angka yang berbeda dari yang tampil di layar tetap dianggap gagal.

Ketiga, skala besar. Ribuan transaksi dicatat, lalu diperiksa apakah
laporan tetap benar, apakah jumlahnya cocok, dan berapa lama prosesnya.

Cara pakai:
    python tools/uji_impor_ekspor_skala.py
"""
from __future__ import annotations

import io
import os
import shutil
import sys
import tempfile
import time
from pathlib import Path

AKAR = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(AKAR / "src"))

FOLDER = Path(tempfile.mkdtemp(prefix="akuntuntas_skala_"))
os.environ["LOCALAPPDATA"] = str(FOLDER)

from akuntansi_id import db, modules, services  # noqa: E402
from akuntansi_id import modules_ops as O  # noqa: E402
from akuntansi_id.core import accounting as acc  # noqa: E402


class Pemeriksa:
    def __init__(self):
        self.lulus = 0
        self.gagal = 0
        self.temuan: list[str] = []

    def cek(self, nama: str, syarat: bool, catatan: str = ""):
        if syarat:
            self.lulus += 1
            print(f"  [BENAR] {nama}")
        else:
            self.gagal += 1
            self.temuan.append(nama)
            print(f"  [SALAH] {nama}")
            if catatan:
                print(f"          {catatan}")

    def cek_sama(self, nama: str, satu, lain, catatan: str = ""):
        self.cek(nama, satu == lain, catatan or f"{satu!r} vs {lain!r}")

    def ringkas(self) -> int:
        print()
        print("=" * 76)
        if self.gagal:
            print(f"  HASIL: {self.lulus} BENAR, {self.gagal} SALAH")
            print()
            for t in self.temuan:
                print(f"    - {t}")
        else:
            print(f"  HASIL: {self.lulus} BENAR, 0 SALAH")
        print("=" * 76)
        return 1 if self.gagal else 0


def main() -> int:
    p = Pemeriksa()
    print("=" * 76)
    print("  UJI IMPOR BERANTAKAN, EKSPOR, DAN SKALA BESAR")
    print("=" * 76)
    print()

    db.init_db()
    b = db.q1("SELECT id FROM companies ORDER BY id LIMIT 1")
    cid = b["id"] if b else services.create_company("PT Skala", bentuk="pt")

    # ==================================================================
    print("[1. Impor jurnal dengan pemisah titik koma]")
    csv_benar = ("tanggal;no_bukti;keterangan;kode_akun;debit;kredit\n"
                 "2026-01-15;BKM-001;Penerimaan jasa;1001;5000000;0\n"
                 "2026-01-15;BKM-001;Pendapatan jasa;4001;0;5000000\n")
    h = O.impor_jurnal_massal(cid, csv_benar)
    p.cek_sama("impor jurnal: 1 entri berhasil", h.get("berhasil"), 1,
               f"hasil: {h}")

    kas = db.scalar("""SELECT COALESCE(SUM(debit - kredit),0) FROM journal_lines
                       WHERE company_id=? AND kode_akun='1001'""", (cid,))
    p.cek_sama("saldo kas setelah impor = Rp5.000.000", int(kas), 5_000_000)
    print()

    # ==================================================================
    print("[2. Impor jurnal dengan angka berformat Indonesia]")
    # Angka 1.500.000,00 dan 2.000.000 ditulis dengan gaya Indonesia.
    csv_angka = ("tanggal;no_bukti;keterangan;kode_akun;debit;kredit\n"
                 "2026-02-10;BKM-002;Penjualan;1001;1.500.000,00;0\n"
                 "2026-02-10;BKM-002;Penjualan;4001;0;1.500.000,00\n")
    h2 = O.impor_jurnal_massal(cid, csv_angka)
    nilai = db.scalar("""SELECT COALESCE(SUM(debit),0) FROM journal_lines
                         WHERE company_id=? AND kode_akun='1001'""", (cid,))
    # Saldo sebelumnya 5.000.000, ditambah 1.500.000.
    p.cek_sama("angka bergaya Indonesia terbaca benar (1.500.000,00)",
               int(nilai), 6_500_000, f"hasil impor: {h2}")
    print()

    # ==================================================================
    print("[3. Impor jurnal dengan berbagai bentuk tanggal]")
    # Berkas dari aplikasi lain memakai bentuk yang berbeda beda.
    bentuk_tanggal = [
        ("15/03/2026", "2026-03-15", "hari di depan (Indonesia)"),
        ("15-03-2026", "2026-03-15", "hari di depan dengan tanda hubung"),
        ("15.03.2026", "2026-03-15", "hari di depan dengan titik"),
        ("2026/03/15", "2026-03-15", "tahun di depan dengan garis miring"),
        ("15 Maret 2026", "2026-03-15", "nama bulan Indonesia"),
        ("15 Mar 2026", "2026-03-15", "nama bulan singkat"),
        ("31/12/2026", "2026-12-31", "akhir tahun"),
        ("01/01/2026", "2026-01-01", "awal tahun"),
        ("29/02/2024", "2024-02-29", "tahun kabisat"),
    ]
    for i, (masuk, harap, nama) in enumerate(bentuk_tanggal, start=1):
        bukti = f"TGL-{i:02d}"
        csv_t = ("tanggal;no_bukti;keterangan;kode_akun;debit;kredit\n"
                 f"{masuk};{bukti};Uji tanggal;1001;1000000;0\n"
                 f"{masuk};{bukti};Uji tanggal;4001;0;1000000\n")
        hh = O.impor_jurnal_massal(cid, csv_t)
        baris_t = db.q1("SELECT tanggal FROM journal_entries WHERE company_id=? AND no_bukti=?",
                        (cid, bukti))
        p.cek(f"tanggal '{masuk}' ({nama}) terbaca sebagai {harap}",
              bool(baris_t) and baris_t["tanggal"] == harap,
              f"tersimpan: {baris_t['tanggal'] if baris_t else 'tidak ada'} "
              f"(impor: {hh.get('berhasil')} berhasil, {hh.get('gagal')} gagal)")

    # Tanggal yang memang tidak sah harus tetap ditolak.
    csv_salah = ("tanggal;no_bukti;keterangan;kode_akun;debit;kredit\n"
                 "45/13/2026;TGL-99;Tanggal salah;1001;1000000;0\n"
                 "45/13/2026;TGL-99;Tanggal salah;4001;0;1000000\n")
    hs = O.impor_jurnal_massal(cid, csv_salah)
    ada_salah = db.scalar("SELECT COUNT(*) FROM journal_entries WHERE company_id=? AND no_bukti='TGL-99'",
                          (cid,))
    p.cek("tanggal tidak sah (45/13/2026) tetap ditolak",
          int(ada_salah) == 0 and hs.get("gagal") == 1, f"hasil: {hs}")
    print()

    # ==================================================================
    print("[4. Impor jurnal dengan tanda kutip dan pemisah di dalam teks]")
    csv_kutip = ('tanggal,no_bukti,keterangan,kode_akun,debit,kredit\n'
                 '2026-04-01,BKM-004,"Penjualan, termasuk PPN",1001,1100000,0\n'
                 '2026-04-01,BKM-004,"Penjualan, termasuk PPN",4001,0,1000000\n'
                 '2026-04-01,BKM-004,"Penjualan, termasuk PPN",2013,0,100000\n')
    h4 = O.impor_jurnal_massal(cid, csv_kutip)
    ket = db.q1("""SELECT keterangan FROM journal_entries
                   WHERE company_id=? AND no_bukti=?""", (cid, "BKM-004"))
    p.cek("keterangan bertanda koma di dalam kutip terbaca utuh",
          bool(ket) and "termasuk PPN" in (ket["keterangan"] or ""),
          f"keterangan: {ket['keterangan']!r}" if ket else "entri tidak ada")
    print()

    # ==================================================================
    print("[5. Impor dengan kode akun tidak dikenal ditolak]")
    csv_akun_palsu = ("tanggal;no_bukti;keterangan;kode_akun;debit;kredit\n"
                      "2026-05-01;BKM-999;Akun palsu;9999;1000000;0\n"
                      "2026-05-01;BKM-999;Akun palsu;4001;0;1000000\n")
    h5 = O.impor_jurnal_massal(cid, csv_akun_palsu)
    ada_palsu = db.scalar("""SELECT COUNT(*) FROM journal_lines
                             WHERE company_id=? AND kode_akun='9999'""", (cid,))
    p.cek_sama("akun 9999 tidak masuk ke basis data", int(ada_palsu), 0)
    p.cek_sama("impor melaporkan kegagalan", h5.get("gagal"), 1, f"hasil: {h5}")
    print()

    # ==================================================================
    print("[6. Impor jurnal tidak seimbang ditolak]")
    csv_tidak_seimbang = ("tanggal;no_bukti;keterangan;kode_akun;debit;kredit\n"
                          "2026-05-02;BKM-998;Tidak seimbang;1001;1000000;0\n"
                          "2026-05-02;BKM-998;Tidak seimbang;4001;0;900000\n")
    h6 = O.impor_jurnal_massal(cid, csv_tidak_seimbang)
    ada = db.scalar("SELECT COUNT(*) FROM journal_entries WHERE company_id=? AND no_bukti='BKM-998'",
                    (cid,))
    p.cek_sama("jurnal tidak seimbang tidak tersimpan", int(ada), 0)
    p.cek_sama("impor melaporkan 1 gagal", h6.get("gagal"), 1, f"hasil: {h6}")
    print()

    # ==================================================================
    print("[7. Impor mitra dengan duplikat dan baris kosong]")
    csv_mitra = ("nama;npwp;email;telepon\n"
                 "PT Satu;01.111.111.1-111.000;a@x.id;0811\n"
                 "\n"
                 "PT Satu;01.111.111.1-111.000;a@x.id;0811\n"
                 "PT Dua;;b@x.id;0822\n"
                 ";;;\n")
    hm = O.impor_mitra_massal(cid, csv_mitra)
    jumlah_mitra = db.scalar("SELECT COUNT(*) FROM partners WHERE company_id=?",
                             (cid,))
    p.cek_sama("impor mitra: 2 mitra tersimpan", int(jumlah_mitra), 2,
               f"hasil: {hm}")
    p.cek_sama("impor mitra: 1 duplikat dikenali", hm.get("duplikat"), 1,
               f"hasil: {hm}")
    print()

    # ==================================================================
    print("[8. Impor produk dengan harga bergaya Indonesia]")
    csv_produk = ("kode;nama;satuan;tipe;harga_beli;harga_jual;stok\n"
                  "P001;Kopi 250g;pcs;barang;45.000;75.000;100\n"
                  "P002;Jasa Konsultasi;jam;jasa;0;350.000;0\n")
    hp = O.impor_produk_massal(cid, csv_produk)
    prod = db.q1("SELECT harga_beli, harga_jual FROM products WHERE company_id=? AND kode='P001'",
                 (cid,))
    if prod:
        p.cek_sama("harga beli 45.000 terbaca Rp45.000",
                   int(prod["harga_beli"]), 45_000)
        p.cek_sama("harga jual 75.000 terbaca Rp75.000",
                   int(prod["harga_jual"]), 75_000)
    else:
        p.cek("produk P001 tersimpan", False, f"hasil: {hp}")
    print()

    # ==================================================================
    print("[9. Impor CSV rusak tidak membuat aplikasi berhenti]")
    for nama, isi in (
        ("berkas kosong", ""),
        ("hanya header", "tanggal;no_bukti;kode_akun;debit;kredit\n"),
        ("tanpa kolom wajib", "a;b;c\n1;2;3\n"),
        ("teks bukan CSV", "ini bukan berkas csv sama sekali\n\n\n"),
        ("baris sangat panjang", "tanggal;no_bukti;keterangan;kode_akun;debit;kredit\n"
                                 "2026-01-01;X;" + "A" * 50_000 + ";1001;1;0\n"),
        ("kolom kurang", "tanggal;no_bukti;kode_akun;debit;kredit\n"
                         "2026-01-01;Y;1001;1\n"),
    ):
        try:
            O.impor_jurnal_massal(cid, isi)
            p.cek(f"CSV rusak ({nama}) tidak membuat aplikasi berhenti", True)
        except ValueError:
            p.cek(f"CSV rusak ({nama}) ditolak dengan keterangan", True)
        except Exception as e:
            p.cek(f"CSV rusak ({nama}) tidak membuat aplikasi berhenti", False,
                  f"{type(e).__name__}: {e}")
    print()

    # ==================================================================
    print("[10. Ekspor Excel: isinya diperiksa]")
    try:
        sys.path.insert(0, str(AKAR / "src"))
        from akuntansi_id.ui.pages.laporan import ekspor_excel
        berkas = FOLDER / "uji.xlsx"
        data_tabel = [
            ("Neraca", [["Akun", "Jumlah"], ["Kas", 6_500_000], ["Modal", 6_500_000]]),
            ("Laba Rugi", [["Akun", "Jumlah"], ["Pendapatan", 3_500_000]]),
        ]
        ekspor_excel("Uji", data_tabel, str(berkas))
        p.cek("berkas Excel terbentuk", berkas.exists(),
              f"ukuran: {berkas.stat().st_size if berkas.exists() else 0} byte")
        if berkas.exists():
            import openpyxl
            wb = openpyxl.load_workbook(str(berkas))
            p.cek_sama("Excel memuat 2 lembar", len(wb.sheetnames), 2,
                       f"lembar: {wb.sheetnames}")
            ws = wb[wb.sheetnames[0]]
            p.cek_sama("angka di Excel sama dengan data (6.500.000)",
                       ws.cell(row=2, column=2).value, 6_500_000)
            p.cek_sama("teks di Excel sama dengan data",
                       ws.cell(row=1, column=1).value, "Akun")
    except RuntimeError as e:
        print(f"          dilewati: {e}")
    except Exception as e:
        p.cek("ekspor Excel berjalan tanpa galat", False,
              f"{type(e).__name__}: {e}")
    print()

    # ==================================================================
    print("[11. Ekspor PDF: isinya diperiksa]")
    try:
        from akuntansi_id.ui.pages.laporan import ekspor_pdf
        berkas_pdf = FOLDER / "uji.pdf"
        ekspor_pdf("Laporan Uji", "PT Skala", data_tabel, str(berkas_pdf))
        p.cek("berkas PDF terbentuk", berkas_pdf.exists(),
              f"ukuran: {berkas_pdf.stat().st_size if berkas_pdf.exists() else 0} byte")
        if berkas_pdf.exists():
            isi = berkas_pdf.read_bytes()
            p.cek("PDF berisi penanda berkas PDF yang sah",
                  isi[:5] == b"%PDF-", f"awal berkas: {isi[:8]!r}")
            p.cek("PDF berukuran wajar (bukan berkas kosong)",
                  len(isi) > 1000, f"ukuran: {len(isi)} byte")
            # Isi PDF dimampatkan oleh pustaka pembuatnya, sehingga teksnya
            # tidak dapat dicari langsung. Karena itu yang diperiksa adalah
            # jumlah halaman dan bahwa berkasnya dapat dibuka kembali.
            jumlah_halaman = isi.count(b"/Type /Page") - isi.count(b"/Type /Pages")
            p.cek("PDF memuat sekurangnya satu halaman", jumlah_halaman >= 1,
                  f"halaman terdeteksi: {jumlah_halaman}")
            # Berkas yang hanya berisi kop tanpa tabel akan jauh lebih kecil.
            p.cek("PDF memuat tabel, bukan hanya kop",
                  len(isi) > 1500, f"ukuran: {len(isi)} byte")
    except RuntimeError as e:
        print(f"          dilewati: {e}")
    except Exception as e:
        p.cek("ekspor PDF berjalan tanpa galat", False,
              f"{type(e).__name__}: {e}")
    print()

    # ==================================================================
    print("[12. Ekspor CSV: isinya diperiksa]")
    try:
        from akuntansi_id.ui.pages.laporan import ekspor_csv
        berkas_csv = FOLDER / "uji.csv"
        ekspor_csv("Uji", [["Akun", "Jumlah"], ["Kas", 6_500_000]], str(berkas_csv))
        teks = berkas_csv.read_text(encoding="utf-8-sig")
        p.cek("CSV memuat baris data", "Kas" in teks and "6500000" in teks,
              f"isi: {teks[:80]!r}")
    except Exception as e:
        p.cek("ekspor CSV berjalan tanpa galat", False,
              f"{type(e).__name__}: {e}")
    print()

    # ==================================================================
    print("[13. Skala besar: 3.000 transaksi]")
    jumlah = 3000
    mulai = time.time()
    kode_akun_pasangan = [("1001", "4001"), ("1001", "4002"),
                          ("1002", "4001"), ("6006", "1001")]
    with db.tx() as conn:
        for i in range(jumlah):
            d, k = kode_akun_pasangan[i % len(kode_akun_pasangan)]
            nilai = (i % 500 + 1) * 10_000
            bulan = (i % 12) + 1
            hari = (i % 27) + 1
            tanggal = f"2026-{bulan:02d}-{hari:02d}"
            cur = conn.execute(
                """INSERT INTO journal_entries(company_id, tanggal, no_bukti,
                   keterangan, sumber)
                   VALUES(?,?,?,?,'uji')""",
                (cid, tanggal, f"SKALA-{i:05d}", f"Transaksi skala {i}"))
            eid = cur.lastrowid
            conn.execute("""INSERT INTO journal_lines(company_id, entry_id,
                            kode_akun, debit, kredit)
                            VALUES(?,?,?,?,0)""", (cid, eid, d, nilai))
            conn.execute("""INSERT INTO journal_lines(company_id, entry_id,
                            kode_akun, debit, kredit)
                            VALUES(?,?,?,0,?)""", (cid, eid, k, nilai))
    lama_catat = time.time() - mulai

    total_jurnal = db.scalar("SELECT COUNT(*) FROM journal_entries WHERE company_id=?",
                             (cid,))
    print(f"          {jumlah} transaksi tercatat dalam {lama_catat:.1f} detik")
    print(f"          total jurnal di basis data: {total_jurnal}")

    # Laporan harus tetap benar pada skala ini.
    mulai = time.time()
    ner = acc.neraca(cid, 2026)
    lr = acc.laba_rugi(cid, 2026)
    lama_lapor = time.time() - mulai

    harta = getattr(ner, "total_aset", 0)
    utang_modal = (getattr(ner, "total_liabilitas", 0)
                   + getattr(ner, "total_ekuitas", 0))
    p.cek_sama("skala besar: harta = utang + ekuitas", harta, utang_modal,
               f"harta {harta:,} vs utang+modal {utang_modal:,}")
    print(f"          laporan dihitung dalam {lama_lapor:.2f} detik")

    # Neraca saldo harus seimbang.
    ns = acc.total_neraca_saldo(cid, 2026)
    p.cek_sama("skala besar: neraca saldo debit = kredit",
               ns.get("debit", 0), ns.get("kredit", 0),
               f"debit {ns.get('debit', 0):,} vs kredit {ns.get('kredit', 0):,}")

    # Setiap entri skala besar punya dua baris. Entri dari pengujian impor
    # sebelumnya dapat punya lebih dari dua baris, jadi yang dihitung
    # hanya entri bersumber 'uji'.
    baris_skala = db.scalar("""SELECT COUNT(*) FROM journal_lines jl
                               JOIN journal_entries je ON je.id = jl.entry_id
                               WHERE jl.company_id=? AND je.sumber='uji'""",
                            (cid,))
    entri_skala = db.scalar("""SELECT COUNT(*) FROM journal_entries
                               WHERE company_id=? AND sumber='uji'""", (cid,))
    p.cek_sama("skala besar: setiap entri punya tepat 2 baris",
               int(baris_skala), int(entri_skala) * 2)

    p.cek("laporan pada skala besar selesai dalam waktu wajar (< 10 detik)",
          lama_lapor < 10, f"lama: {lama_lapor:.2f} detik")
    print()

    # ==================================================================
    print("[14. Skala besar: ringkasan bulanan cocok dengan setahun]")
    ringkas = acc.ringkasan_bulanan(cid, 2026)
    if ringkas:
        jumlah_bulanan = sum(r.get("pendapatan", 0) for r in ringkas)
        pendapatan_setahun = (getattr(lr, "pendapatan_usaha", 0)
                              + getattr(lr, "pendapatan_lain", 0))
        p.cek_sama("jumlah pendapatan bulanan = pendapatan setahun",
                   jumlah_bulanan, pendapatan_setahun)
        p.cek_sama("jumlah baris ringkasan = 12 bulan", len(ringkas), 12)
    else:
        p.cek("ringkasan bulanan terisi pada skala besar", False)
    print()

    # ==================================================================
    print("[15. Skala besar: buku besar cocok dengan saldo]")
    # Hitungan pembanding harus menyaring tahun yang sama dengan buku besar,
    # karena ada data uji dari tahun lain (misalnya uji tahun kabisat).
    for kode in ("1001", "4001", "6006"):
        bk = acc.buku_besar(cid, kode, 2026)
        if bk:
            akhir = bk[-1].get("saldo")
            r = db.q1("""SELECT COALESCE(SUM(jl.debit),0) AS d,
                                COALESCE(SUM(jl.kredit),0) AS k
                         FROM journal_lines jl
                         JOIN journal_entries je ON je.id = jl.entry_id
                         WHERE jl.company_id=? AND jl.kode_akun=?
                           AND je.tanggal BETWEEN '2026-01-01' AND '2026-12-31'""",
                      (cid, kode))
            a = db.q1("SELECT normal, saldo_awal FROM accounts WHERE company_id=? AND kode=?",
                      (cid, kode))
            normal = (a["normal"] or "Debit").lower() if a else "debit"
            net = r["d"] - r["k"]
            # Saldo awal ikut diperhitungkan bila ada.
            awal = int(a["saldo_awal"] or 0) if a else 0
            harap = (awal - net) if normal.startswith("kredit") else (awal + net)
            p.cek_sama(f"skala besar: buku besar {kode} cocok", akhir, harap,
                       f"aplikasi {akhir:,} vs hitungan {harap:,}")
    print()

    return p.ringkas()


if __name__ == "__main__":
    try:
        kode = main()
    finally:
        shutil.rmtree(FOLDER, ignore_errors=True)
    sys.exit(kode)
