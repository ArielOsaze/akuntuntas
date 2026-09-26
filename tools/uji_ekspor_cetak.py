"""
Uji ekspor Excel dan cetak A4.

Yang diperiksa bukan hanya apakah berkasnya terbentuk, tetapi apakah
isinya benar: angka tersimpan sebagai angka, lembar terpisah per tabel,
ukuran kertas benar benar A4, dan hasil cetak memuat seluruh baris.

Cara pakai:
    python tools/uji_ekspor_cetak.py
"""
from __future__ import annotations

import os
import shutil
import sys
import tempfile
from pathlib import Path

AKAR = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(AKAR / "src"))

FOLDER = Path(tempfile.mkdtemp(prefix="akuntuntas_ekspor_"))
os.environ["LOCALAPPDATA"] = str(FOLDER)
# Qt tidak menampilkan jendela pada pengujian ini.
os.environ.setdefault("QT_QPA_PLATFORM", "offscreen")

from PySide6.QtWidgets import QApplication  # noqa: E402

app = QApplication.instance() or QApplication([])

from akuntansi_id.ui import ekspor_cetak as EC  # noqa: E402


class Pemeriksa:
    def __init__(self):
        self.lulus = 0
        self.gagal = 0
        self.temuan: list[str] = []

    def cek(self, nama: str, syarat: bool, catatan: str = ""):
        if syarat:
            self.lulus += 1
            print(f"  [BENAR] {nama}")
        else:
            self.gagal += 1
            self.temuan.append(nama)
            print(f"  [SALAH] {nama}")
            if catatan:
                print(f"          {catatan}")

    def cek_sama(self, nama: str, satu, lain, catatan: str = ""):
        self.cek(nama, satu == lain,
                 catatan or f"seharusnya {lain!r}, ternyata {satu!r}")

    def ringkas(self) -> int:
        print()
        print("=" * 76)
        if self.gagal:
            print(f"  HASIL: {self.lulus} BENAR, {self.gagal} SALAH")
            print()
            for t in self.temuan:
                print(f"    - {t}")
        else:
            print(f"  HASIL: {self.lulus} BENAR, 0 SALAH")
        print("=" * 76)
        return 1 if self.gagal else 0


def main() -> int:
    p = Pemeriksa()
    print("=" * 76)
    print("  UJI EKSPOR EXCEL DAN CETAK A4")
    print("=" * 76)
    print()

    # Data uji menyerupai laporan sungguhan.
    data = [
        ("Neraca", [
            ["Akun", "Kode", "Jumlah"],
            ["Kas", "1001", 750_000_000],
            ["Piutang Usaha", "1101", 250_000_000],
            ["Persediaan", "1104", 180_000_000],
            ["Total Harta", "", 1_180_000_000],
        ]),
        ("Laba Rugi", [
            ["Keterangan", "Jumlah"],
            ["Pendapatan", 375_000_000],
            ["Harga Pokok Penjualan", 63_600_000],
            ["Laba Kotor", 311_400_000],
        ]),
    ]

    # ==================================================================
    print("[1. Ekspor Excel: berkas dan isinya]")
    berkas = FOLDER / "laporan.xlsx"
    try:
        EC.ekspor_excel("Laporan Keuangan", data, str(berkas),
                        subjudul="Tahun 2026",
                        nama_perusahaan="PT Sinar Nusantara")
        p.cek("berkas Excel terbentuk", berkas.exists(),
              f"ukuran: {berkas.stat().st_size if berkas.exists() else 0} byte")
    except Exception as e:
        p.cek("ekspor Excel berjalan tanpa galat", False,
              f"{type(e).__name__}: {e}")
        berkas = None

    if berkas and berkas.exists():
        import openpyxl
        wb = openpyxl.load_workbook(str(berkas))

        p.cek_sama("satu lembar per tabel", len(wb.sheetnames), 2,
                   f"lembar: {wb.sheetnames}")
        p.cek("nama lembar sama dengan nama tabel",
              wb.sheetnames[0] == "Neraca", f"lembar: {wb.sheetnames}")

        ws = wb["Neraca"]
        p.cek_sama("judul ada di kop", ws.cell(row=1, column=1).value,
                   "Laporan Keuangan")
        # Kop memuat perusahaan dan waktu.
        kop = str(ws.cell(row=2, column=1).value or "")
        p.cek("kop memuat nama perusahaan", "Sinar Nusantara" in kop,
              f"kop: {kop}")
        p.cek("kop memuat subjudul", "2026" in kop, f"kop: {kop}")

        # Baris judul tabel di baris 5, data mulai baris 6.
        p.cek_sama("baris judul tabel benar",
                   ws.cell(row=5, column=1).value, "Akun")
        p.cek_sama("data pertama benar",
                   ws.cell(row=6, column=1).value, "Kas")

        # Angka harus tersimpan sebagai angka, bukan teks.
        sel_angka = ws.cell(row=6, column=3)
        p.cek("angka tersimpan sebagai angka, bukan teks",
              isinstance(sel_angka.value, (int, float)),
              f"tipe: {type(sel_angka.value).__name__}, "
              f"nilai: {sel_angka.value!r}")
        p.cek_sama("nilai angka benar", sel_angka.value, 750_000_000)
        p.cek("angka memakai format ribuan",
              sel_angka.number_format == '#,##0',
              f"format: {sel_angka.number_format}")

        # Teks tetap teks.
        sel_teks = ws.cell(row=6, column=1)
        p.cek("teks tersimpan sebagai teks",
              isinstance(sel_teks.value, str))

        # Baris judul dibekukan.
        p.cek("baris judul dibekukan", ws.freeze_panes is not None,
              f"freeze_panes: {ws.freeze_panes}")

        # Lebar kolom sudah diatur.
        lebar = ws.column_dimensions["A"].width
        p.cek("lebar kolom sudah diatur", lebar and lebar >= 12,
              f"lebar kolom A: {lebar}")

        # Jumlah baris data benar.
        terisi = [r for r in range(6, 12) if ws.cell(row=r, column=1).value]
        p.cek_sama("seluruh baris data tertulis", len(terisi), 4,
                   f"baris terisi: {terisi}")

    # ==================================================================
    print()
    print("[2. Nama lembar yang bermasalah dibersihkan]")
    sulit = [
        ("Neraca 2026/2027", "karakter garis miring"),
        ("Laba: Rugi", "karakter titik dua"),
        ("Data [penting]", "karakter kurung siku"),
        ("A" * 50, "nama terlalu panjang"),
        ("", "nama kosong"),
        ("'Kutip'", "tanda kutip"),
    ]
    for nama, alasan in sulit:
        bersih = EC.nama_lembar_aman(nama)
        p.cek(f"nama lembar aman ({alasan})",
              len(bersih) <= 31
              and not any(c in bersih for c in ':\\/?*[]')
              and bool(bersih),
              f"{nama!r} -> {bersih!r}")

    # Nama yang sama dua kali tidak boleh bertabrakan.
    dipakai: set = set()
    n1 = EC.nama_lembar_aman("Neraca", dipakai)
    n2 = EC.nama_lembar_aman("Neraca", dipakai)
    p.cek("nama lembar kembar dibedakan", n1 != n2, f"{n1!r} vs {n2!r}")

    # ==================================================================
    print()
    print("[3. Cetak A4: ukuran kertas dan isinya]")
    pdf = FOLDER / "cetak.pdf"
    try:
        EC.simpan_pdf(None, "Laporan Keuangan", "Tahun 2026", data, str(pdf))
        p.cek("berkas PDF terbentuk", pdf.exists(),
              f"ukuran: {pdf.stat().st_size if pdf.exists() else 0} byte")
    except Exception as e:
        p.cek("penyimpanan PDF berjalan tanpa galat", False,
              f"{type(e).__name__}: {e}")
        pdf = None

    if pdf and pdf.exists():
        isi = pdf.read_bytes()
        p.cek("penanda berkas PDF sah", isi[:5] == b"%PDF-",
              f"awal: {isi[:8]!r}")

        # Ukuran halaman A4 dalam satuan titik PDF: 595 x 842.
        import re
        ukuran = re.search(rb"/MediaBox\s*\[\s*([\d.]+)\s+([\d.]+)\s+"
                           rb"([\d.]+)\s+([\d.]+)", isi)
        if ukuran:
            lebar = float(ukuran.group(3)) - float(ukuran.group(1))
            tinggi = float(ukuran.group(4)) - float(ukuran.group(2))
            # A4 = 595 x 842 titik. Toleransi 2 titik untuk pembulatan.
            p.cek("ukuran halaman A4 tegak (595 x 842 titik)",
                  abs(lebar - 595) <= 2 and abs(tinggi - 842) <= 2,
                  f"terukur: {lebar:.0f} x {tinggi:.0f} titik")
        else:
            p.cek("ukuran halaman dapat dibaca", False,
                  "penanda MediaBox tidak ditemukan")

        p.cek("PDF memuat sekurangnya satu halaman",
              isi.count(b"/Type /Page") >= 1)
        p.cek("PDF berukuran wajar", len(isi) > 1500,
              f"ukuran: {len(isi)} byte")

    # ==================================================================
    print()
    print("[4. Cetak A4 mendatar]")
    pdf_l = FOLDER / "cetak_mendatar.pdf"
    try:
        EC.simpan_pdf(None, "Daftar Panjang", "", data, str(pdf_l),
                      mendatar=True)
        isi = pdf_l.read_bytes()
        import re
        ukuran = re.search(rb"/MediaBox\s*\[\s*([\d.]+)\s+([\d.]+)\s+"
                           rb"([\d.]+)\s+([\d.]+)", isi)
        if ukuran:
            lebar = float(ukuran.group(3)) - float(ukuran.group(1))
            tinggi = float(ukuran.group(4)) - float(ukuran.group(2))
            p.cek("halaman mendatar (842 x 595 titik)",
                  abs(lebar - 842) <= 2 and abs(tinggi - 595) <= 2,
                  f"terukur: {lebar:.0f} x {tinggi:.0f} titik")
        else:
            p.cek("ukuran halaman mendatar dapat dibaca", False)
    except Exception as e:
        p.cek("penyimpanan PDF mendatar berjalan", False,
              f"{type(e).__name__}: {e}")

    # ==================================================================
    print()
    print("[5. Tabel panjang terpecah menjadi beberapa halaman]")
    # 120 baris data, jauh lebih panjang dari satu halaman A4.
    panjang = [("Daftar Panjang", [["No", "Nama", "Jumlah"]]
                + [[i, f"Baris {i}", i * 1_000_000] for i in range(1, 121)])]
    pdf_p = FOLDER / "panjang.pdf"
    try:
        EC.simpan_pdf(None, "Uji Halaman", "", panjang, str(pdf_p))
        isi = pdf_p.read_bytes()
        halaman = isi.count(b"/Type /Page") - isi.count(b"/Type /Pages")
        p.cek("tabel panjang terpecah menjadi beberapa halaman",
              halaman >= 2, f"jumlah halaman: {halaman}")
        p.cek("berkas tidak kosong", len(isi) > 3000,
              f"ukuran: {len(isi)} byte")
    except Exception as e:
        p.cek("tabel panjang dapat dicetak", False,
              f"{type(e).__name__}: {e}")

    # ==================================================================
    print()
    print("[6. Pratinjau gambar terbentuk]")
    try:
        gambar = EC.pratinjau_gambar("Pratinjau", "Uji", data, lebar=800)
        p.cek("gambar pratinjau terbentuk", not gambar.isNull(),
              f"ukuran: {gambar.width()}x{gambar.height()}")
        p.cek("gambar pratinjau berukuran wajar", gambar.height() > 100,
              f"tinggi: {gambar.height()}")
        # Gambar harus berisi sesuatu, bukan putih kosong.
        berwarna = 0
        for y in range(0, gambar.height(), 20):
            for x in range(0, gambar.width(), 20):
                if gambar.pixel(x, y) != 0xFFFFFFFF:
                    berwarna += 1
        p.cek("gambar pratinjau memuat isi, bukan halaman kosong",
              berwarna > 10, f"titik berisi: {berwarna}")
        jalur_gambar = FOLDER / "pratinjau.png"
        gambar.save(str(jalur_gambar))
        p.cek("pratinjau dapat disimpan sebagai gambar",
              jalur_gambar.exists())
    except Exception as e:
        p.cek("pratinjau gambar berjalan tanpa galat", False,
              f"{type(e).__name__}: {e}")

    # ==================================================================
    print()
    print("[7. Nilai kosong dan teks panjang tetap aman]")
    aneh = [
        ("Data Kosong", [["A", "B"]]),
        ("Nilai Kosong", [["A", "B"], [None, ""], ["x", None]]),
        ("Teks Panjang", [["Keterangan"], ["A" * 500]]),
        ("Karakter Khusus", [["Nama"], ["<b>tebal</b> & 'kutip' \"ganda\""]]),
        ("Angka Nol", [["Jumlah"], [0], [0.0]]),
    ]
    berkas2 = FOLDER / "aneh.xlsx"
    try:
        EC.ekspor_excel("Uji Nilai Aneh", aneh, str(berkas2))
        p.cek("ekspor dengan nilai kosong dan karakter khusus berhasil",
              berkas2.exists())
        import openpyxl
        wb = openpyxl.load_workbook(str(berkas2))
        p.cek_sama("seluruh lembar terbentuk", len(wb.sheetnames), 5,
                   f"lembar: {wb.sheetnames}")
    except Exception as e:
        p.cek("ekspor nilai aneh berjalan tanpa galat", False,
              f"{type(e).__name__}: {e}")

    pdf_aneh = FOLDER / "aneh.pdf"
    try:
        EC.simpan_pdf(None, "Uji Nilai Aneh", "", aneh, str(pdf_aneh))
        p.cek("cetak dengan nilai kosong dan karakter khusus berhasil",
              pdf_aneh.exists())
        # Karakter khusus tidak boleh merusak tata letak HTML.
        isi = pdf_aneh.read_bytes()
        p.cek("berkas hasil tetap sah", isi[:5] == b"%PDF-")
    except Exception as e:
        p.cek("cetak nilai aneh berjalan tanpa galat", False,
              f"{type(e).__name__}: {e}")

    # ==================================================================
    print()
    print("[8. Teks tabel diubah menjadi angka dengan benar]")
    # Angka harus dikenali supaya dapat dijumlahkan di Excel, sedangkan
    # teks biasa tetap tersimpan sebagai teks.
    contoh = [
        ("750.000.000", 750_000_000, "angka ribuan gaya Indonesia"),
        ("Rp1.500.000", 1_500_000, "angka dengan awalan Rp"),
        ("45.000,50", 45_000.5, "angka dengan desimal koma"),
        ("-2.500.000", -2_500_000, "angka negatif"),
        ("12,5%", 0.125, "persentase"),
        ("0", 0, "nol"),
        ("Kas", None, "teks biasa"),
        ("PT Maju Bersama", None, "nama perusahaan"),
        ("", None, "teks kosong"),
        ("2026-01-15", None, "tanggal"),
        ("INV/2026/0001", None, "nomor dokumen"),
    ]
    for teks, harap, alasan in contoh:
        hasil = EC._teks_ke_angka(teks)
        p.cek_sama(f"ubah teks ke angka ({alasan}): {teks!r}", hasil, harap)

    # ==================================================================
    print()
    print("[9. Mengambil isi tabel di layar]")
    from PySide6.QtWidgets import QTableWidget, QTableWidgetItem

    tabel = QTableWidget(2, 3)
    tabel.setHorizontalHeaderLabels(["Nama", "Jumlah", "Catatan"])
    isi_tabel = [
        ("Kas", "750.000.000", "saldo awal"),
        ("Piutang", "250.000.000", ""),
    ]
    for r, baris in enumerate(isi_tabel):
        for c, nilai in enumerate(baris):
            tabel.setItem(r, c, QTableWidgetItem(nilai))

    diambil = EC.tabel_dari_widget(tabel)
    p.cek_sama("jumlah baris termasuk judul", len(diambil), 3)
    p.cek_sama("baris judul diambil dari kepala tabel", diambil[0],
               ["Nama", "Jumlah", "Catatan"])
    p.cek_sama("teks tetap teks", diambil[1][0], "Kas")
    p.cek_sama("angka diubah menjadi angka", diambil[1][1], 750_000_000)
    p.cek("angka bertipe angka, bukan teks",
          isinstance(diambil[1][1], (int, float)),
          f"tipe: {type(diambil[1][1]).__name__}")

    # Tabel kosong tidak membuat galat.
    p.cek_sama("tabel kosong menghasilkan daftar kosong",
               EC.tabel_dari_widget(None), [])

    return p.ringkas()


if __name__ == "__main__":
    try:
        kode = main()
    finally:
        shutil.rmtree(FOLDER, ignore_errors=True)
    sys.exit(kode)
